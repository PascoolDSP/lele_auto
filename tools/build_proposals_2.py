"""Second batch of stock-sheet proposals (11-20), inspired by the template research of 2026-09-25.

Same brief as batch 1 (products, suppliers, 4 weeks per month, stock / ordered / received,
automatic totals); each file explores a different layout pattern found in existing templates.
"""
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from build_proposals import (
    EXAMPLE_FONT, FONT, GREEN_FILL, INPUT_FILL, N, ORANGE_FILL, PRODUCTS, RED_FILL, SUPPLIERS,
    WEEK_DATA, Theme, calc, cell, example, guide, header, inp, last_filled, negative_red,
    non_blank_sum, nonneg_validation, product_columns, product_ref, products_sheet, save,
    simple_month, title, widths,
)

LEFT = Alignment(horizontal="left", vertical="center")
GREY_FILL = PatternFill("solid", fgColor="E7E6E6")
BLUE_FILL = PatternFill("solid", fgColor="CFE2F3")
ZONES = {"Farine T55": "Réserve sèche", "Farine T45": "Réserve sèche", "Sucre semoule": "Réserve sèche",
         "Sucre glace": "Réserve sèche", "Beurre doux": "Chambre froide", "Crème liquide 35%": "Chambre froide",
         "Lait entier": "Chambre froide", "Oeufs": "Chambre froide", "Chocolat blanc": "Réserve sèche",
         "Chocolat noir 70%": "Réserve sèche", "Levure chimique": "Réserve sèche", "Vanille (gousse)": "Réserve sèche"}
ESSENTIALS = {"Farine T55", "Sucre semoule", "Beurre doux", "Crème liquide 35%", "Oeufs", "Chocolat blanc"}


def dv_list(ws, options, rng):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)


def cf_text(ws, rng, first_ref, text, fill):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first_ref}="{text}"'], fill=fill))


def left_input(ws, ref, value):
    c = inp(ws, ref, value, example=True)
    c.alignment = LEFT
    return c


def product_names():
    return [p[0] for p in PRODUCTS]


def week_block_headers(ws, t, col, labels, week_label, header_row=4, date=True):
    end = col + len(labels) - 1
    header(ws, f"{L(col)}{header_row}", week_label, t, span_to=f"{L(end)}{header_row}" if end > col else None)
    if date:
        ws.merge_cells(f"{L(col)}{header_row + 1}:{L(end)}{header_row + 1}")
        inp(ws, f"{L(col)}{header_row + 1}", None, fmt="DD/MM/YYYY")
        for row in ws[f"{L(col)}{header_row + 1}:{L(end)}{header_row + 1}"]:
            for c in row:
                c.fill = INPUT_FILL
                c.border = BORDER
    sub = header_row + (2 if date else 1)
    for j, lab in enumerate(labels):
        header(ws, f"{L(col + j)}{sub}", lab, t, soft=True)
        ws.column_dimensions[L(col + j)].width = 10.5
    return {lab: L(col + j) for j, lab in enumerate(labels)}


THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ================================================================ 11
def p11():
    t = Theme("4A6FA5")
    wb = Workbook()
    guide(wb, t, "11 · Dans l'ordre de la réserve",
          "Tes produits sont rangés dans l'ordre où tu les vois quand tu fais le tour : réserve sèche, chambre froide, "
          "congélateur. Le vendredi, tu fais ton tour en descendant la page, sans jamais chercher une ligne.",
          ["1. Dans l'onglet du mois, écris tes produits dans leur zone, dans l'ordre de tes étagères.",
           "2. Le vendredi : fais ton tour et remplis Stock puis Commandé, zone par zone.",
           "3. À la livraison : Reçu. Les totaux du mois se font tout seuls, l'écart devient rouge s'il manque quelque chose.",
           "4. Nouveau mois : duplique l'onglet et efface les chiffres (garde les noms de produits)."],
          ["Octobre 2026 — tes produits groupés par zone de rangement, 4 semaines côte à côte."])
    ws = wb.create_sheet("Octobre 2026")
    ws.sheet_view.showGridLines = False
    title(ws, "Octobre 2026", "Produits dans l'ordre de ta réserve · Stock / Commandé / Reçu chaque semaine", t)
    header(ws, "A4", "Produit", t, span_to="A6")
    header(ws, "B4", "Unité", t, span_to="B6")
    widths(ws, {"A": 24, "B": 8})
    cols = {}
    col = 3
    for w in range(4):
        cols[w] = week_block_headers(ws, t, col, ["Stock", "Commandé", "Reçu"], f"Semaine {w + 1}")
        col += 3
    tot = {lab: L(col + j) for j, lab in enumerate(("Commandé", "Reçu", "Écart"))}
    header(ws, f"{L(col)}4", "Total du mois", t, span_to=f"{L(col + 2)}5")
    for lab, letter in tot.items():
        header(ws, f"{letter}6", lab, t, soft=True)
        ws.column_dimensions[letter].width = 11
    zones = [("Réserve sèche", 16), ("Chambre froide", 12), ("Congélateur", 6)]
    r = 7
    idx = {p[0]: i for i, p in enumerate(PRODUCTS)}
    for zone, size in zones:
        ws.merge_cells(f"A{r}:{tot['Écart']}{r}")
        cell(ws, f"A{r}", zone.upper(), bold=True, color="FFFFFF", fill=t.fill)
        ws.row_dimensions[r].height = 20
        r += 1
        items = [p for p in PRODUCTS if ZONES[p[0]] == zone]
        start = r
        for k in range(size):
            p = items[k] if k < len(items) else None
            left_input(ws, f"A{r}", p[0] if p else None)
            inp(ws, f"B{r}", p[2] if p else None, example=True)
            for w in range(4):
                ex = example(w, idx[p[0]]) if p else (None, None, None)
                for lab, v in zip(("Stock", "Commandé", "Reçu"), ex):
                    inp(ws, f"{cols[w][lab]}{r}", v, example=True)
            calc(ws, f"{tot['Commandé']}{r}", non_blank_sum([f"{cols[w]['Commandé']}{r}" for w in range(4)]), bold=True)
            calc(ws, f"{tot['Reçu']}{r}", non_blank_sum([f"{cols[w]['Reçu']}{r}" for w in range(4)]), bold=True)
            a, b = f"{tot['Commandé']}{r}", f"{tot['Reçu']}{r}"
            calc(ws, f"{tot['Écart']}{r}", f'=IF(OR({a}="",{b}=""),"",{b}-{a})', bold=True)
            r += 1
        nonneg_validation(ws, f"C{start}:{L(col - 1)}{r - 1}")
        e = tot["Écart"]
        negative_red(ws, f"{e}{start}:{e}{r - 1}", f"{e}{start}")
    ws.freeze_panes = "C7"
    save(wb, "11_ordre_de_la_reserve.xlsx")


# ================================================================ 12
def p12():
    t = Theme("8C5E3C")
    wb = Workbook()
    dates = 8
    guide(wb, t, "12 · Le cadencier",
          "L'outil classique des cuisines françaises : pour chaque fournisseur, une ligne par produit et une colonne "
          "par commande. Tu vois d'un coup d'œil ce que tu as commandé les fois précédentes, ta commande habituelle "
          "et combien de fois il a manqué quelque chose.",
          ["1. Chaque onglet = un fournisseur. Écris ses produits dans la colonne Produit.",
           "2. Le jour de la commande : tape la date en haut de la colonne suivante, puis Stock et Commandé.",
           "3. À la livraison : Reçu, dans la même colonne.",
           "4. À droite, Commande habituelle et Dernière commande t'aident à décider sans tout recalculer."],
          ["Metro, Crèmerie, Transgourmet — un cadencier par fournisseur (8 commandes par page)."])
    for sname, _, tel, jc, jl in SUPPLIERS:
        ws = wb.create_sheet(sname)
        ws.sheet_view.showGridLines = False
        title(ws, f"Cadencier — {sname}", f"Commande le {jc} · livraison le {jl} · Tél. {tel}", t)
        header(ws, "A4", "Produit", t, span_to="A6")
        header(ws, "B4", "Unité", t, span_to="B6")
        widths(ws, {"A": 24, "B": 8})
        cols = {}
        col = 3
        for d in range(dates):
            cols[d] = week_block_headers(ws, t, col, ["Stock", "Cmd", "Reçu"], f"Commande {d + 1}")
            for lab in ("Stock", "Cmd", "Reçu"):
                ws.column_dimensions[cols[d][lab]].width = 8
            col += 3
        for d, day in ((0, 2), (1, 9)):
            c = ws[f"{cols[d]['Stock']}5"]
            c.value = f"=DATE(2026,10,{day})"
            c.font = EXAMPLE_FONT
        right = {lab: L(col + j) for j, lab in enumerate(("Commande habituelle", "Dernière commande", "Fois où il a manqué"))}
        header(ws, f"{L(col)}4", "Pour t'aider", t, span_to=f"{L(col + 2)}5")
        for lab, letter in right.items():
            header(ws, f"{letter}6", lab, t, soft=True)
            ws.column_dimensions[letter].width = 13
        items = [(i, p) for i, p in enumerate(PRODUCTS) if p[1] == sname]
        for k in range(20):
            r = 7 + k
            gi, p = items[k] if k < len(items) else (None, None)
            left_input(ws, f"A{r}", p[0] if p else None)
            inp(ws, f"B{r}", p[2] if p else None, example=True)
            for d in range(dates):
                ex = example(d, gi) if gi is not None else (None, None, None)
                for lab, v in zip(("Stock", "Cmd", "Reçu"), ex):
                    inp(ws, f"{cols[d][lab]}{r}", v, example=True)
            cmds = [f"{cols[d]['Cmd']}{r}" for d in range(dates)]
            calc(ws, f"{right['Commande habituelle']}{r}",
                 f'=IF(COUNT({",".join(cmds)})=0,"",ROUND(AVERAGE({",".join(cmds)}),1))', bold=True)
            calc(ws, f"{right['Dernière commande']}{r}", last_filled(cmds))
            pairs = "+".join(f'AND(ISNUMBER({cols[d]["Cmd"]}{r}),ISNUMBER({cols[d]["Reçu"]}{r}),'
                             f'N({cols[d]["Reçu"]}{r})<N({cols[d]["Cmd"]}{r}))*1' for d in range(dates))
            calc(ws, f"{right['Fois où il a manqué']}{r}", f'=IF(A{r}="","",{pairs})')
        m = right["Fois où il a manqué"]
        ws.conditional_formatting.add(f"{m}7:{m}26", FormulaRule(formula=[f"AND(ISNUMBER({m}7),{m}7>0)"], fill=RED_FILL))
        for d in range(dates):
            c_, r_ = cols[d]["Cmd"], cols[d]["Reçu"]
            ws.conditional_formatting.add(
                f"{r_}7:{r_}26", FormulaRule(formula=[f"AND(ISNUMBER({r_}7),ISNUMBER({c_}7),{r_}7<{c_}7)"], fill=RED_FILL))
        nonneg_validation(ws, f"C7:{L(col - 1)}26")
        ws.freeze_panes = "C7"
    save(wb, "12_cadencier.xlsx")


# ================================================================ 13
def p13():
    t = Theme("2B7A78")
    wb = Workbook()
    guide(wb, t, "13 · La liste à cocher",
          "Une page par semaine, comme une liste de courses. Le vendredi tu remplis ton stock et ce que tu commandes, "
          "puis tu coches « Commandé ». Le lundi, tu coches « Reçu » ligne par ligne en vérifiant les cartons : "
          "la ligne se barre et devient grise. Ce qui reste non barré n'est pas arrivé.",
          ["1. Remplis l'onglet Produits une seule fois.",
           "2. Vendredi : Stock, Quantité à commander, puis choisis ✓ dans « Commandé ? ».",
           "3. Lundi : pour chaque produit livré, tape la quantité reçue et choisis ✓ dans « Reçu ? ».",
           "4. Rouge = reçu moins que commandé. Le récap du mois additionne les 4 semaines."],
          ["Produits — ta liste.", "Semaine 1 à 4 — une liste à cocher par semaine.", "Récap du mois — les totaux."])
    products_sheet(wb, t)
    for w in range(4):
        ws = wb.create_sheet(f"Semaine {w + 1}")
        ws.sheet_view.showGridLines = False
        title(ws, f"Semaine {w + 1}", "Coche ✓ au fur et à mesure — la ligne se barre quand c'est reçu", t)
        heads = [("Produit", 24), ("Fournisseur", 14), ("Unité", 7), ("Stock", 9), ("À commander", 14),
                 ("Commandé ?", 13), ("Reçu", 9), ("Reçu ?", 9), ("Écart", 8)]
        for i, (h, wd) in enumerate(heads, start=1):
            header(ws, f"{L(i)}4", h, t)
            ws.column_dimensions[L(i)].width = wd
        product_columns(ws, 5, t)
        for i in range(N):
            r = 5 + i
            ex = example(w, i)
            inp(ws, f"D{r}", ex[0], example=True)
            inp(ws, f"E{r}", ex[1], example=True)
            inp(ws, f"F{r}", ("✓" if ex[1] else None) if ex[1] is not None else None, example=True)
            inp(ws, f"G{r}", ex[2], example=True)
            ticked = ex[1] and (w > 0 or i < 7)  # week 1 example: delivery check still in progress
            inp(ws, f"H{r}", "✓" if ticked else None, example=True)
            calc(ws, f"I{r}", f'=IF(OR(E{r}="",G{r}=""),"",G{r}-E{r})')
        last = 4 + N
        dv_list(ws, ["✓"], f"F5:F{last}")
        dv_list(ws, ["✓"], f"H5:H{last}")
        nonneg_validation(ws, f"D5:E{last}")
        nonneg_validation(ws, f"G5:G{last}")
        negative_red(ws, f"I5:I{last}", "I5")  # added first: a shortage stays red even once ticked
        done = DifferentialStyle(font=Font(strike=True, color="8C8C8C"), fill=GREY_FILL)
        ws.conditional_formatting.add(f"A5:I{last}", Rule(type="expression", dxf=done, formula=[f'$H5="✓"']))
        ws.freeze_panes = "B5"
    rc = wb.create_sheet("Récap du mois")
    rc.sheet_view.showGridLines = False
    title(rc, "Récap du mois", "Calculé à partir des 4 semaines — rien à remplir", t)
    heads = [("Produit", 24), ("Unité", 8), ("Total commandé", 14), ("Total reçu", 12), ("Manque", 10), ("Dernier stock", 13)]
    for i, (h, wd) in enumerate(heads, start=1):
        header(rc, f"{L(i)}4", h, t)
        rc.column_dimensions[L(i)].width = wd
    for i in range(N):
        r, s = 5 + i, 5 + i
        cell(rc, f"A{r}", product_ref("A", i), bold=True)
        cell(rc, f"B{r}", product_ref("C", i), color="555555", align="center")
        calc(rc, f"C{r}", non_blank_sum([f"'Semaine {w + 1}'!E{s}" for w in range(4)]), bold=True)
        calc(rc, f"D{r}", non_blank_sum([f"'Semaine {w + 1}'!G{s}" for w in range(4)]), bold=True)
        calc(rc, f"E{r}", f'=IF(OR(C{r}="",D{r}=""),"",D{r}-C{r})')
        calc(rc, f"F{r}", last_filled([f"'Semaine {w + 1}'!D{s}" for w in range(4)]), bold=True)
    negative_red(rc, f"E5:E{4 + N}", "E5")
    save(wb, "13_liste_a_cocher.xlsx")


# ================================================================ 14
def p14():
    t = Theme("444444")
    wb = Workbook()
    guide(wb, t, "14 · La fiche papier",
          "Pour celles qui préfèrent le stylo en réserve : une fiche A4 à imprimer, avec tes produits déjà écrits et "
          "de grandes cases vides. Tu la remplis à la main, puis tu recopies les chiffres dans l'onglet du mois quand "
          "tu as deux minutes. Les calculs se font là.",
          ["1. Remplis l'onglet Produits une seule fois.",
           "2. Chaque vendredi : imprime l'onglet « Fiche à imprimer » (il tient sur une page A4).",
           "3. Remplis-la au stylo pendant ton tour de réserve, puis à la livraison.",
           "4. Recopie les chiffres dans l'onglet du mois : les totaux et les écarts se calculent."],
          ["Produits — ta liste.", "Fiche à imprimer — la fiche papier de la semaine.",
           "Octobre 2026 — la saisie et les calculs du mois."])
    products_sheet(wb, t)
    ws = wb.create_sheet("Fiche à imprimer")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "FICHE DE STOCK — SEMAINE DU ____ / ____ / ________"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = "Compté par : ______________________          Livraison reçue le : ____ / ____"
    ws["A2"].font = Font(name=FONT, size=10, color="444444")
    heads = [("Produit", 30), ("Unité", 8), ("Stock compté", 14), ("À commander", 14), ("Reçu", 12), ("OK ?", 7), ("Remarque", 26)]
    dark = Side(style="thin", color="555555")
    box = Border(left=dark, right=dark, top=dark, bottom=dark)
    for i, (h, wd) in enumerate(heads, start=1):
        c = cell(ws, f"{L(i)}4", h, bold=True, color="FFFFFF", fill=t.fill, align="center")
        c.border = box
        ws.column_dimensions[L(i)].width = wd
    for i in range(30):
        r = 5 + i
        c = cell(ws, f"A{r}", product_ref("A", i), size=11)
        c.border = box
        c = cell(ws, f"B{r}", product_ref("C", i), align="center", size=11)
        c.border = box
        for col in "CDEFG":
            cell(ws, f"{col}{r}").border = box
        ws.row_dimensions[r].height = 22
    ws.print_title_rows = "4:4"
    ws.print_area = "A1:G34"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    simple_month(wb, t, "Octobre 2026", True)
    save(wb, "14_fiche_papier.xlsx")


# ================================================================ 15
def p15():
    t = Theme("B03A2E")
    wb = Workbook()
    guide(wb, t, "15 · Les dates de péremption",
          "Le suivi du stock habituel, plus un onglet pour les produits frais : à chaque livraison de crème, beurre, "
          "lait ou œufs, tu notes la date limite. Le fichier compte les jours restants et colore : rouge = périmé ou "
          "à utiliser aujourd'hui, orange = dans les 3 jours.",
          ["1. Remplis l'onglet Produits une seule fois.",
           "2. Chaque semaine : Stock, Commandé, Reçu dans l'onglet du mois, comme d'habitude.",
           "3. À la livraison des produits frais : une ligne dans « Dates limites » avec la DLC écrite sur l'emballage.",
           "4. Quand un lot est fini, choisis Oui dans « Terminé ? » : la ligne devient grise."],
          ["Produits — ta liste.", "Octobre 2026 — stock, commandé, reçu.", "Dates limites — les lots frais et leur DLC."])
    products_sheet(wb, t)
    simple_month(wb, t, "Octobre 2026", True)
    ws = wb.create_sheet("Dates limites")
    ws.sheet_view.showGridLines = False
    title(ws, "Dates limites (DLC)", "Jours restants calculés à partir de la date du jour", t)
    heads = [("Reçu le", 12), ("Produit", 24), ("Quantité", 10), ("DLC", 12), ("Jours restants", 13), ("État", 14), ("Terminé ?", 11)]
    for i, (h, wd) in enumerate(heads, start=1):
        header(ws, f"{L(i)}4", h, t)
        ws.column_dimensions[L(i)].width = wd
    examples = [("=TODAY()-3", "Crème liquide 35%", 6, "=TODAY()+1", "Non"),
                ("=TODAY()-3", "Oeufs", 90, "=TODAY()+12", "Non"),
                ("=TODAY()-3", "Lait entier", 6, "=TODAY()+3", "Non"),
                ("=TODAY()-3", "Beurre doux", 5, "=TODAY()+20", "Non"),
                ("=TODAY()-10", "Crème liquide 35%", 6, "=TODAY()-2", "Oui")]
    rows = 80
    for k in range(rows):
        r = 5 + k
        ex = examples[k] if k < len(examples) else (None,) * 5
        inp(ws, f"A{r}", ex[0], example=True, fmt="DD/MM/YYYY")
        left_input(ws, f"B{r}", ex[1])
        inp(ws, f"C{r}", ex[2], example=True)
        inp(ws, f"D{r}", ex[3], example=True, fmt="DD/MM/YYYY")
        calc(ws, f"E{r}", f'=IF(D{r}="","",D{r}-TODAY())')
        calc(ws, f"F{r}", f'=IF(E{r}="","",IF(G{r}="Oui","Terminé",IF(E{r}<0,"PÉRIMÉ",IF(E{r}<=1,"AUJOURD\'HUI",IF(E{r}<=3,"BIENTÔT","OK")))))')
        inp(ws, f"G{r}", ex[4], example=True)
    last = 4 + rows
    dv = DataValidation(type="list", formula1=f"=Produits!$A$5:$A${4 + N}", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B5:B{last}")
    dv_list(ws, ["Oui", "Non"], f"G5:G{last}")
    grey = DifferentialStyle(font=Font(color="8C8C8C"), fill=GREY_FILL)
    ws.conditional_formatting.add(f"A5:G{last}", Rule(type="expression", dxf=grey, formula=['$G5="Oui"']))
    for text, fill in (("PÉRIMÉ", RED_FILL), ("AUJOURD'HUI", RED_FILL), ("BIENTÔT", ORANGE_FILL), ("OK", GREEN_FILL)):
        cf_text(ws, f"F5:F{last}", "F5", text.replace('"', ''), fill)
    ws.freeze_panes = "A5"
    save(wb, "15_dates_de_peremption.xlsx")


# ================================================================ 16
def p16():
    t = Theme("7B4B94")
    wb = Workbook()
    reasons = ["Périmé", "Raté", "Cassé / renversé", "Dégustation", "Autre"]
    guide(wb, t, "16 · Les pertes",
          "Le suivi du stock habituel, plus un carnet des pertes : produit périmé, préparation ratée, pot renversé, "
          "dégustation. À la fin du mois, tu vois ce qui part à la poubelle et pourquoi — c'est souvent ce qui "
          "explique qu'il te manque de la marchandise.",
          ["1. Remplis l'onglet Produits une seule fois.",
           "2. Chaque semaine : Stock, Commandé, Reçu dans l'onglet du mois.",
           "3. Quand tu jettes ou perds quelque chose : une ligne dans « Pertes » (date, produit, quantité, raison).",
           "4. « Bilan des pertes » : tape le mois et l'année, le total par produit et par raison s'affiche."],
          ["Produits — ta liste.", "Octobre 2026 — stock, commandé, reçu.", "Pertes — le carnet.",
           "Bilan des pertes — par produit et par raison."])
    products_sheet(wb, t)
    simple_month(wb, t, "Octobre 2026", True)
    ws = wb.create_sheet("Pertes")
    ws.sheet_view.showGridLines = False
    title(ws, "Carnet des pertes", "Une ligne par perte", t)
    heads = [("Date", 12), ("Produit", 24), ("Quantité", 10), ("Unité", 8), ("Raison", 18), ("Note", 30), ("Mois", 7), ("Année", 7)]
    for i, (h, wd) in enumerate(heads, start=1):
        header(ws, f"{L(i)}4", h, t)
        ws.column_dimensions[L(i)].width = wd
    examples = [("=DATE(2026,10,3)", "Crème liquide 35%", 1, "Périmé", "Bouteille ouverte oubliée"),
                ("=DATE(2026,10,6)", "Chocolat blanc", 0.3, "Raté", "Ganache tranchée"),
                ("=DATE(2026,10,8)", "Oeufs", 6, "Cassé / renversé", "Plateau tombé"),
                ("=DATE(2026,10,9)", "Beurre doux", 0.2, "Dégustation", "Essai nouvelle recette")]
    rows = 150
    for k in range(rows):
        r = 5 + k
        ex = examples[k] if k < len(examples) else (None,) * 5
        inp(ws, f"A{r}", ex[0], example=True, fmt="DD/MM/YYYY")
        left_input(ws, f"B{r}", ex[1])
        inp(ws, f"C{r}", ex[2], example=True)
        calc(ws, f"D{r}", f'=IF(B{r}="","",IFERROR(INDEX(Produits!$C$5:$C${4 + N},MATCH(B{r},Produits!$A$5:$A${4 + N},0)),""))')
        left_input(ws, f"E{r}", ex[3])
        left_input(ws, f"F{r}", ex[4])
        calc(ws, f"G{r}", f'=IF(A{r}="","",MONTH(A{r}))')
        calc(ws, f"H{r}", f'=IF(A{r}="","",YEAR(A{r}))')
    last = 4 + rows
    dv = DataValidation(type="list", formula1=f"=Produits!$A$5:$A${4 + N}", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B5:B{last}")
    dv_list(ws, reasons, f"E5:E{last}")
    nonneg_validation(ws, f"C5:C{last}")
    ws.freeze_panes = "A5"
    bl = wb.create_sheet("Bilan des pertes")
    bl.sheet_view.showGridLines = False
    title(bl, "Bilan des pertes", "Tape le mois et l'année dans les cases jaunes", t)
    cell(bl, "A3", "Mois (1 à 12)", bold=True)
    inp(bl, "B3", 10)
    cell(bl, "C3", "Année", bold=True)
    inp(bl, "D3", 2026)
    heads = [("Produit", 24), ("Unité", 8)] + [(r_, 13) for r_ in reasons] + [("Total perdu", 12)]
    for i, (h, wd) in enumerate(heads, start=1):
        header(bl, f"{L(i)}5", h, t)
        bl.column_dimensions[L(i)].width = wd
    rng = lambda c: f"Pertes!${c}$5:${c}${last}"
    for i in range(N):
        r = 6 + i
        cell(bl, f"A{r}", product_ref("A", i), bold=True)
        cell(bl, f"B{r}", product_ref("C", i), color="555555", align="center")
        for j, reason in enumerate(reasons):
            calc(bl, f"{L(3 + j)}{r}",
                 f'=IF($A{r}="","",IF(SUMIFS({rng("C")},{rng("B")},$A{r},{rng("E")},"{reason}",{rng("G")},$B$3,{rng("H")},$D$3)=0,"",'
                 f'SUMIFS({rng("C")},{rng("B")},$A{r},{rng("E")},"{reason}",{rng("G")},$B$3,{rng("H")},$D$3)))')
        tl = L(3 + len(reasons))
        calc(bl, f"{tl}{r}", f'=IF(COUNT(C{r}:{L(2 + len(reasons))}{r})=0,"",SUM(C{r}:{L(2 + len(reasons))}{r}))', bold=True)
    tl = L(3 + len(reasons))
    bl.conditional_formatting.add(f"{tl}6:{tl}{5 + N}", FormulaRule(formula=[f"ISNUMBER({tl}6)"], fill=ORANGE_FILL))
    bl.freeze_panes = "B6"
    save(wb, "16_pertes.xlsx")


# ================================================================ 17
def p17():
    t = Theme("1D6FA5")
    wb = Workbook()
    guide(wb, t, "17 · Les tendances en barres",
          "Le tableau du mois habituel, avec à droite une petite barre par semaine qui montre comment ton stock "
          "évolue, et une flèche : ▲ il remonte, ▼ il baisse. D'un coup d'œil, même sur téléphone, tu vois les "
          "produits qui fondent.",
          ["1. Remplis l'onglet Produits une seule fois.",
           "2. Chaque semaine : Stock, Commandé, Reçu.",
           "3. Les barres et la flèche se dessinent toutes seules au fil des semaines.",
           "4. Nouveau mois : duplique l'onglet du mois."],
          ["Produits — ta liste.", "Octobre 2026 — la saisie, les totaux et les tendances."])
    products_sheet(wb, t)
    ws, first, tot = simple_month(wb, t, "Octobre 2026", True)
    stock_cols = [L(4 + 3 * w) for w in range(4)]
    base = 4 + 12 + 4
    header(ws, f"{L(base)}4", "Tendance du stock", t, span_to=f"{L(base + 4)}5")
    for j, lab in enumerate(("S1", "S2", "S3", "S4", "")):
        header(ws, f"{L(base + j)}6", lab or "Évol.", t, soft=True)
        ws.column_dimensions[L(base + j)].width = 7 if j < 4 else 7
    for i in range(N):
        r = first + i
        refs = [f"{c}{r}" for c in stock_cols]
        mx = f"MAX({','.join(refs)})"
        for w, ref in enumerate(refs):
            c = calc(ws, f"{L(base + w)}{r}",
                     f'=IF(OR({ref}="",{mx}=0),"",REPT("█",MAX(1,ROUND({ref}/{mx}*6,0))))')
            c.font = Font(name=FONT, size=8, color=t.main)
            c.alignment = Alignment(horizontal="left", vertical="bottom")
        filled = f'COUNT({",".join(refs)})'
        first_v = f'IF({refs[0]}<>"",{refs[0]},IF({refs[1]}<>"",{refs[1]},IF({refs[2]}<>"",{refs[2]},{refs[3]})))'
        last_v = last_filled(refs)[1:]
        calc(ws, f"{L(base + 4)}{r}",
             f'=IF({filled}<2,"",IF({last_v}>{first_v},"▲",IF({last_v}<{first_v},"▼","=")))', bold=True)
    ev = L(base + 4)
    last = first + N - 1
    ws.conditional_formatting.add(f"{ev}{first}:{ev}{last}", FormulaRule(formula=[f'{ev}{first}="▼"'], font=Font(color="C0392B"), fill=RED_FILL))
    ws.conditional_formatting.add(f"{ev}{first}:{ev}{last}", FormulaRule(formula=[f'{ev}{first}="▲"'], font=Font(color="1E7B34"), fill=GREEN_FILL))
    save(wb, "17_tendances_en_barres.xlsx")


# ================================================================ 18
def p18():
    t = Theme("C0762B")
    wb = Workbook()
    guide(wb, t, "18 · L'essentiel chaque semaine",
          "Tu ne recomptes pas tout chaque vendredi : seulement tes produits essentiels (farine, sucre, beurre, "
          "crème, œufs, chocolat blanc…), ceux qui partent vite. Le reste, tu le comptes une fois par mois. "
          "Ton comptage du vendredi prend deux minutes.",
          ["1. Onglet Produits : mets Oui dans « Essentiel ? » pour les produits à suivre chaque semaine.",
           "2. Chaque vendredi, onglet « Chaque semaine » : ils apparaissent tout seuls. Stock, Commandé, Reçu.",
           "3. Une fois par mois, onglet « Une fois par mois » : tous les autres produits, stock et commande.",
           "4. Pour changer un produit de liste, change juste Oui / Non dans Produits."],
          ["Produits — ta liste, avec la colonne Essentiel.", "Chaque semaine — les essentiels, 4 semaines.",
           "Une fois par mois — tous les autres produits."])
    ws = wb.create_sheet("Produits")
    ws.sheet_view.showGridLines = False
    title(ws, "Mes produits", "Oui = compté chaque semaine · Non = compté une fois par mois", t)
    heads = [("Produit", 26), ("Fournisseur", 18), ("Unité", 9), ("Essentiel ?", 12), ("N° essentiel", 12), ("N° autre", 10)]
    for i, (h, wd) in enumerate(heads, start=1):
        header(ws, f"{L(i)}4", h, t, soft=i > 4)
        ws.column_dimensions[L(i)].width = wd
    for i in range(N):
        r = 5 + i
        p = PRODUCTS[i] if i < len(PRODUCTS) else None
        left_input(ws, f"A{r}", p[0] if p else None)
        left_input(ws, f"B{r}", p[1] if p else None)
        inp(ws, f"C{r}", p[2] if p else None, example=True)
        inp(ws, f"D{r}", ("Oui" if p[0] in ESSENTIALS else "Non") if p else None, example=True)
        c = calc(ws, f"E{r}", f'=IF(AND(A{r}<>"",D{r}="Oui"),SUMPRODUCT(($A$5:A{r}<>"")*($D$5:D{r}="Oui")),"")')
        c.font = Font(name=FONT, color="999999", size=9)
        c = calc(ws, f"F{r}", f'=IF(AND(A{r}<>"",D{r}<>"Oui"),SUMPRODUCT(($A$5:A{r}<>"")*($D$5:D{r}<>"Oui")),"")')
        c.font = Font(name=FONT, color="999999", size=9)
    dv_list(ws, ["Oui", "Non"], f"D5:D{4 + N}")
    ws.freeze_panes = "A5"

    def pick(k, col, rank_col):
        return (f'=IFERROR(INDEX(Produits!${col}$5:${col}${4 + N},'
                f'MATCH({k},Produits!${rank_col}$5:${rank_col}${4 + N},0)),"")')

    wk = wb.create_sheet("Chaque semaine")
    wk.sheet_view.showGridLines = False
    title(wk, "Chaque semaine — mes essentiels", "La liste se remplit toute seule à partir des Oui de l'onglet Produits", t)
    header(wk, "A4", "Produit", t, span_to="A6")
    header(wk, "B4", "Unité", t, span_to="B6")
    widths(wk, {"A": 24, "B": 8})
    cols = {}
    col = 3
    for w in range(4):
        cols[w] = week_block_headers(wk, t, col, ["Stock", "Commandé", "Reçu"], f"Semaine {w + 1}")
        col += 3
    tot = {lab: L(col + j) for j, lab in enumerate(("Commandé", "Reçu", "Écart"))}
    header(wk, f"{L(col)}4", "Total du mois", t, span_to=f"{L(col + 2)}5")
    for lab, letter in tot.items():
        header(wk, f"{letter}6", lab, t, soft=True)
        wk.column_dimensions[letter].width = 11
    ess = [i for i, p in enumerate(PRODUCTS) if p[0] in ESSENTIALS]
    for k in range(20):
        r = 7 + k
        cell(wk, f"A{r}", pick(k + 1, "A", "E"), bold=True)
        cell(wk, f"B{r}", pick(k + 1, "C", "E"), color="555555", align="center")
        for w in range(4):
            ex = example(w, ess[k]) if k < len(ess) else (None, None, None)
            for lab, v in zip(("Stock", "Commandé", "Reçu"), ex):
                inp(wk, f"{cols[w][lab]}{r}", v, example=True)
        calc(wk, f"{tot['Commandé']}{r}", non_blank_sum([f"{cols[w]['Commandé']}{r}" for w in range(4)]), bold=True)
        calc(wk, f"{tot['Reçu']}{r}", non_blank_sum([f"{cols[w]['Reçu']}{r}" for w in range(4)]), bold=True)
        a, b = f"{tot['Commandé']}{r}", f"{tot['Reçu']}{r}"
        calc(wk, f"{tot['Écart']}{r}", f'=IF(OR({a}="",{b}=""),"",{b}-{a})', bold=True)
    e = tot["Écart"]
    negative_red(wk, f"{e}7:{e}26", f"{e}7")
    nonneg_validation(wk, f"C7:{L(col - 1)}26")
    wk.freeze_panes = "C7"

    mo = wb.create_sheet("Une fois par mois")
    mo.sheet_view.showGridLines = False
    title(mo, "Une fois par mois — les autres produits", "Compte-les en fin de mois", t)
    cell(mo, "A3", "Date du comptage :", bold=True, border=False)
    inp(mo, "B3", "=DATE(2026,10,30)", example=True, fmt="DD/MM/YYYY")
    heads = [("Produit", 24), ("Unité", 8), ("Stock compté", 13), ("Commandé", 11), ("Reçu", 10), ("Écart", 9)]
    for i, (h, wd) in enumerate(heads, start=1):
        header(mo, f"{L(i)}5", h, t)
        mo.column_dimensions[L(i)].width = wd
    others = [i for i, p in enumerate(PRODUCTS) if p[0] not in ESSENTIALS]
    for k in range(30):
        r = 6 + k
        cell(mo, f"A{r}", pick(k + 1, "A", "F"), bold=True)
        cell(mo, f"B{r}", pick(k + 1, "C", "F"), color="555555", align="center")
        ex = example(0, others[k]) if k < len(others) else (None, None, None)
        for col_, v in zip("CDE", ex):
            inp(mo, f"{col_}{r}", v, example=True)
        calc(mo, f"F{r}", f'=IF(OR(D{r}="",E{r}=""),"",E{r}-D{r})')
    negative_red(mo, "F6:F35", "F6")
    nonneg_validation(mo, "C6:E35")
    mo.freeze_panes = "B6"
    save(wb, "18_essentiel_chaque_semaine.xlsx")


# ================================================================ 19
def p19():
    t = Theme("3E7D5A")
    wb = Workbook()
    guide(wb, t, "19 · Ce que j'ai utilisé",
          "À partir de tes comptages, le fichier calcule ce que tu as utilisé chaque semaine : ce que tu avais, "
          "plus ce que tu as reçu, moins ce qu'il te reste la semaine suivante. Aucune recette à saisir. "
          "Au bout d'un mois, tu sais combien de farine ou de crème tu utilises vraiment par semaine.",
          ["1. Remplis l'onglet Produits une seule fois.",
           "2. Chaque semaine : Stock, Commandé, Reçu.",
           "3. « Utilisé » se calcule dès que tu as compté la semaine suivante.",
           "4. Pour la semaine 4, tape le premier comptage du mois suivant dans la dernière colonne jaune."],
          ["Produits — ta liste.", "Octobre 2026 — comptages, commandes, et ce que tu as utilisé."])
    products_sheet(wb, t)
    ws = wb.create_sheet("Octobre 2026")
    ws.sheet_view.showGridLines = False
    title(ws, "Octobre 2026", "Utilisé = stock de la semaine + reçu − stock de la semaine suivante", t)
    for c, h, w in (("A", "Produit", 22), ("B", "Fournisseur", 14), ("C", "Unité", 7)):
        header(ws, f"{c}4", h, t, span_to=f"{c}6")
        ws.column_dimensions[c].width = w
    product_columns(ws, 7, t)
    labels = ["Stock", "Commandé", "Reçu", "Utilisé"]
    cols = {}
    col = 4
    for w in range(4):
        cols[w] = week_block_headers(ws, t, col, labels, f"Semaine {w + 1}")
        col += 4
    nxt = L(col)
    header(ws, f"{nxt}4", "Mois suivant", t, span_to=f"{nxt}5")
    header(ws, f"{nxt}6", "1er stock", t, soft=True)
    ws.column_dimensions[nxt].width = 11
    tot = {lab: L(col + 1 + j) for j, lab in enumerate(("Utilisé ce mois", "Moyenne / semaine"))}
    header(ws, f"{L(col + 1)}4", "Bilan", t, span_to=f"{L(col + 2)}5")
    for lab, letter in tot.items():
        header(ws, f"{letter}6", lab, t, soft=True)
        ws.column_dimensions[letter].width = 14
    next_ex = [4, 2, 6, 1, 4, 6, 3, 40, 1, 2, 0.5, 6]
    for i in range(N):
        r = 7 + i
        for w in range(4):
            ex = example(w, i)
            for lab, v in zip(("Stock", "Commandé", "Reçu"), ex):
                inp(ws, f"{cols[w][lab]}{r}", v, example=True)
        inp(ws, f"{nxt}{r}", None)
        for w in range(4):
            s, rc = f"{cols[w]['Stock']}{r}", f"{cols[w]['Reçu']}{r}"
            later = f"{cols[w + 1]['Stock']}{r}" if w < 3 else f"{nxt}{r}"
            calc(ws, f"{cols[w]['Utilisé']}{r}",
                 f'=IF(OR({s}="",{later}=""),"",{s}+N({rc})-{later})', bold=True)
            ws[f"{cols[w]['Utilisé']}{r}"].fill = t.softer
        used = [f"{cols[w]['Utilisé']}{r}" for w in range(4)]
        calc(ws, f"{tot['Utilisé ce mois']}{r}", non_blank_sum(used), bold=True)
        calc(ws, f"{tot['Moyenne / semaine']}{r}",
             f'=IF(COUNT({",".join(used)})=0,"",ROUND(AVERAGE({",".join(used)}),1))', bold=True)
    last = 6 + N
    for w in range(4):
        u = cols[w]["Utilisé"]
        ws.conditional_formatting.add(f"{u}7:{u}{last}", FormulaRule(formula=[f"AND(ISNUMBER({u}7),{u}7<0)"], fill=ORANGE_FILL))
    nonneg_validation(ws, f"{nxt}7:{nxt}{last}")
    ws[f"{nxt}3"] = "Orange = négatif : erreur de comptage probable"
    ws[f"{nxt}3"].font = Font(name=FONT, italic=True, size=8, color="8A6D00")
    ws.freeze_panes = "D7"
    save(wb, "19_ce_que_j_ai_utilise.xlsx")


# ================================================================ 20
def p20():
    t = Theme("5B5EA6")
    wb = Workbook()
    statuses = ["À commander", "Commandé", "Livré", "Manquant", "Rien à commander"]
    guide(wb, t, "20 · Le suivi par étiquettes",
          "Chaque produit a une étiquette de couleur par semaine qui dit où il en est : À commander, Commandé, "
          "Livré, Manquant, Rien à commander. En haut, un compteur t'indique combien de produits restent à traiter "
          "cette semaine. Comme un tableau de suivi, mais dans ton tableur.",
          ["1. Remplis l'onglet Produits une seule fois.",
           "2. Vendredi : Stock, puis choisis l'étiquette (À commander, puis Commandé une fois la commande passée).",
           "3. Livraison : Reçu, puis l'étiquette Livré ou Manquant.",
           "4. Le compteur « À traiter » en haut de chaque semaine doit tomber à 0."],
          ["Produits — ta liste.", "Octobre 2026 — stock, commandé, reçu et étiquette pour chaque semaine."])
    products_sheet(wb, t)
    ws = wb.create_sheet("Octobre 2026")
    ws.sheet_view.showGridLines = False
    title(ws, "Octobre 2026", "Étiquette : À commander → Commandé → Livré (ou Manquant)", t)
    for c, h, w in (("A", "Produit", 22), ("B", "Fournisseur", 14), ("C", "Unité", 7)):
        header(ws, f"{c}4", h, t, span_to=f"{c}6")
        ws.column_dimensions[c].width = w
    product_columns(ws, 7, t)
    labels = ["Stock", "Commandé", "Reçu", "Étiquette"]
    cols = {}
    col = 4
    for w in range(4):
        cols[w] = week_block_headers(ws, t, col, labels, f"Semaine {w + 1}")
        ws.column_dimensions[cols[w]["Étiquette"]].width = 20
        col += 4
    last = 6 + N
    cell(ws, "A3", "À traiter :", bold=True, border=False, align="right")
    for w in range(4):
        e = cols[w]["Étiquette"]
        c = calc(ws, f"{e}3", f'=COUNTIF({e}7:{e}{last},"À commander")+COUNTIF({e}7:{e}{last},"Commandé")+COUNTIF({e}7:{e}{last},"Manquant")', bold=True)
        ws.conditional_formatting.add(f"{e}3", FormulaRule(formula=[f"{e}3>0"], fill=ORANGE_FILL))
        ws.conditional_formatting.add(f"{e}3", FormulaRule(formula=[f"{e}3=0"], fill=GREEN_FILL))
    s3 = ["À commander", "Commandé", "À commander", "Rien à commander", "Commandé", "Commandé",
          "À commander", "Commandé", "À commander", "Rien à commander", "Rien à commander", "Rien à commander"]
    for i in range(N):
        r = 7 + i
        for w in range(4):
            ex = example(w, i)
            for lab, v in zip(("Stock", "Commandé", "Reçu"), ex):
                inp(ws, f"{cols[w][lab]}{r}", v, example=True)
            status = None
            if i < len(PRODUCTS):
                if w < 2:
                    s, cm, rc = ex
                    status = "Rien à commander" if cm == 0 else ("Manquant" if rc < cm else "Livré")
                elif w == 2:
                    status = s3[i]
            c = inp(ws, f"{cols[w]['Étiquette']}{r}", status, example=True)
            c.font = Font(name=FONT, bold=True, color="1F1F1F")
    chips = {"À commander": "FBDDB0", "Commandé": "CFE2F3", "Livré": "D3EBD0", "Manquant": "F8C9C4", "Rien à commander": "E7E6E6"}
    for w in range(4):
        e = cols[w]["Étiquette"]
        dv_list(ws, statuses, f"{e}7:{e}{last}")
        for text, color in chips.items():
            cf_text(ws, f"{e}7:{e}{last}", f"{e}7", text, PatternFill("solid", fgColor=color))
        nonneg_validation(ws, f"{cols[w]['Stock']}7:{cols[w]['Reçu']}{last}")
    ws.freeze_panes = "D7"
    save(wb, "20_suivi_par_etiquettes.xlsx")


if __name__ == "__main__":
    for fn in (p11, p12, p13, p14, p15, p16, p17, p18, p19, p20):
        fn()
