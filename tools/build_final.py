"""Merged workbook for Lele (v1), built from her feedback of 2026-10-03 / 10-10.

Base: proposal 8 (12 monthly tabs + yearly view). Added: products grouped by storage zone
(dry store, positive cold room, negative cold room, housekeeping), an order date per week,
a dashboard (spend, incomplete deliveries, low stock, trend bars, top 10 chart) with a month
selector, a suppliers tab and a reception register (date, product, supplier, quantity, lot
number, best-before date) replacing her paper notebook.
"""
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule, Rule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from build_proposals import (
    BORDER, EXAMPLE_FONT, FONT, GREEN_FILL, INPUT_FILL, MONTHS, ORANGE_FILL, RED_FILL, Theme,
    calc, cell, guide, header, inp, last_filled, negative_red, non_blank_sum, nonneg_validation,
    save, title, widths,
)

T = Theme("2E6B3F")
GREY_FILL = PatternFill("solid", fgColor="E7E6E6")
LEFT = Alignment(horizontal="left", vertical="center")
ZONES = ["Réserve sèche", "Chambre froide positive", "Chambre froide négative", "Ménager"]
PER_ZONE = 15
UNITS = ["kg", "g", "L", "pièce", "paquet", "boîte", "sac", "rouleau"]
MONTH_SHORT = ["Janv.", "Févr.", "Mars", "Avr.", "Mai", "Juin", "Juil.", "Août", "Sept.", "Oct.", "Nov.", "Déc."]
SUPPLIERS = [("A et M", "", "", "vendredi", "lundi"), ("Promo Cash", "", "", "vendredi", "lundi"),
             ("Auchan", "", "", "", "")]
# zone -> (name, usual supplier, unit, min stock, unit price EUR) — example data, shown in blue
EXAMPLES = {
    "Réserve sèche": [("Farine T55", "Promo Cash", "kg", 5, 0.9), ("Farine T45", "Promo Cash", "kg", 3, 1.1),
                      ("Sucre semoule", "Promo Cash", "kg", 5, 1.2), ("Sucre roux", "Promo Cash", "kg", 2, 1.8),
                      ("Cannelle", "Promo Cash", "kg", 0.2, 12), ("Vanille liquide", "A et M", "L", 0.5, 25),
                      ("Sel", "Promo Cash", "kg", 1, 0.8), ("Maïzena", "Promo Cash", "kg", 1, 2.5),
                      ("Kinder Bueno", "Promo Cash", "paquet", 2, 4), ("Pistache", "A et M", "kg", 1, 28),
                      ("Nutella", "A et M", "kg", 2, 6), ("Pâte de pistache", "A et M", "kg", 0.5, 35)],
    "Chambre froide positive": [("Crème liquide 35%", "A et M", "L", 6, 4.2), ("Beurre doux", "Promo Cash", "kg", 4, 8.5),
                                ("Oeufs", "Auchan", "pièce", 60, 0.25), ("Lait entier", "Promo Cash", "L", 5, 1.1)],
    "Chambre froide négative": [("Pâte feuilletée surgelée", "Promo Cash", "kg", 2, 5),
                                ("Framboises surgelées", "Promo Cash", "kg", 1, 9)],
    "Ménager": [("Sacs poubelle", "Promo Cash", "rouleau", 2, 3), ("Liquide vaisselle", "Promo Cash", "L", 1, 2.5),
                ("Papier cuisson", "Promo Cash", "rouleau", 2, 4)],
}
SHORT = {("Octobre", 0, "Farine T55"): 7, ("Septembre", 1, "Crème liquide 35%"): 9}  # received < ordered
WEEK_DATES = {"Octobre": [2, 9], "Septembre": [4, 11]}
REGISTER_EXAMPLES = [  # from her notebook photo (14/09 and 18/09)
    ("=DATE(2026,9,14)", "Crème liquide 35%", "A et M", 6, "T079A10968", "=DATE(2026,12,9)"),
    ("=DATE(2026,9,14)", "Oeufs", "Auchan", 90, "FR 22284.002", None),
    ("=DATE(2026,9,14)", "Farine T45", "Promo Cash", 5, "1-223", "=DATE(2027,2,3)"),
    ("=DATE(2026,9,14)", "Vanille liquide", "A et M", 1, "22526", "=DATE(2028,8,18)"),
    ("=DATE(2026,9,14)", "Kinder Bueno", "Promo Cash", 4, "L103219", "=DATE(2027,4,13)"),
    ("=DATE(2026,9,18)", "Pistache", "A et M", 1, "2608609", "=DATE(2027,3,31)"),
    ("=DATE(2026,9,18)", "Pâte de pistache", "A et M", 1, "01-2047 85992", "=DATE(2027,3,31)"),
    ("=DATE(2026,10,5)", "Crème liquide 35%", "A et M", 6, "T081B20011", "=TODAY()+2"),
]


def prod_rows():
    """Yield (zone_index, k, products_row, month_row) for every product slot."""
    for z in range(len(ZONES)):
        for k in range(PER_ZONE):
            yield z, k, 6 + z * (PER_ZONE + 1) + k, 7 + z * (PER_ZONE + 1) + k


def zone_header_rows(sheet):  # 'products' or 'month'
    base = 5 if sheet == "products" else 6
    return [base + z * (PER_ZONE + 1) for z in range(len(ZONES))]


ALL = list(prod_rows())
LAST_PR = ALL[-1][2]
LAST_MR = ALL[-1][3]
PRODUCT_LIST = f"Produits!$G$5:$G${4 + len(ALL)}"
UNIT_LIST = f"Produits!$H$5:$H${4 + len(ALL)}"


def example_product(z, k):
    items = EXAMPLES[ZONES[z]]
    return items[k] if k < len(items) else None


def week_values(month, w, product):
    """Example (stock, ordered, received) for a product in a given month/week, or blanks."""
    if month not in WEEK_DATES or w >= len(WEEK_DATES[month]) or product is None:
        return None, None, None
    name, _, _, mini, _ = product
    stock = round(mini * (1.6 if w == 0 else 1.1), 1)
    stock = int(stock) if float(stock).is_integer() else stock
    cmd = int(mini * 2) if mini >= 1 else round(mini * 2, 1)
    rec = SHORT.get((month, w, name), cmd)
    return stock, cmd, rec


def zone_bar(ws, r, last_col, label):
    ws.merge_cells(f"A{r}:{last_col}{r}")
    cell(ws, f"A{r}", label.upper(), bold=True, color="FFFFFF", fill=T.fill)
    for c in range(2, 27):
        if L(c) <= last_col or len(last_col) > 1:
            ws[f"{L(c)}{r}"].fill = T.fill
    ws.row_dimensions[r].height = 20


# ---------------------------------------------------------------- Produits
def sheet_products(wb):
    ws = wb.create_sheet("Produits")
    ws.sheet_view.showGridLines = False
    title(ws, "Mes produits", "Une ligne par produit, dans sa zone de rangement. Tout le fichier se remplit à partir d'ici.", T)
    heads = [("Produit", 26), ("Fournisseur habituel", 21), ("Unité", 9), ("Stock minimum", 15), ("Prix unitaire (€)", 16)]
    for i, (h, w) in enumerate(heads, start=1):
        header(ws, f"{L(i)}4", h, T)
        ws.column_dimensions[L(i)].width = w
    for z, r in enumerate(zone_header_rows("products")):
        zone_bar(ws, r, "E", ZONES[z])
    for z, k, pr, _ in ALL:
        p = example_product(z, k)
        c = inp(ws, f"A{pr}", p[0] if p else None, example=True); c.alignment = LEFT
        c = inp(ws, f"B{pr}", p[1] if p else None, example=True); c.alignment = LEFT
        inp(ws, f"C{pr}", p[2] if p else None, example=True)
        inp(ws, f"D{pr}", p[3] if p else None, example=True)
        inp(ws, f"E{pr}", p[4] if p else None, example=True, fmt='#,##0.00 "€"')
    # contiguous helper lists (hidden): all product names / units, for dropdowns and lookups
    header(ws, "G4", "Liste", T, soft=True)
    header(ws, "H4", "Unité", T, soft=True)
    for n, (z, k, pr, _) in enumerate(ALL):
        calc(ws, f"G{5 + n}", f'=IF(A{pr}="","",A{pr})').font = Font(name=FONT, color="999999", size=9)
        calc(ws, f"H{5 + n}", f'=IF(A{pr}="","",C{pr})').font = Font(name=FONT, color="999999", size=9)
    ws.column_dimensions["G"].hidden = True
    ws.column_dimensions["H"].hidden = True
    unit_dv = DataValidation(type="list", formula1='"' + ",".join(UNITS) + '"', allow_blank=True)
    ws.add_data_validation(unit_dv)
    sup_dv = DataValidation(type="list", formula1="=Fournisseurs!$A$5:$A$14", allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(sup_dv)
    for z, k, pr, _ in ALL:
        unit_dv.add(f"C{pr}")
        sup_dv.add(f"B{pr}")
    nonneg_validation(ws, f"D6:E{LAST_PR}")
    ws["A3"] = "Prix unitaire : à écrire une seule fois, le tableau de bord le réutilise chaque mois. À changer seulement si le fournisseur change son tarif."
    ws["A3"].font = Font(name=FONT, italic=True, size=9, color="666666")
    ws.freeze_panes = "A5"
    return ws


# ---------------------------------------------------------------- Fournisseurs
def sheet_suppliers(wb):
    ws = wb.create_sheet("Fournisseurs")
    ws.sheet_view.showGridLines = False
    title(ws, "Mes fournisseurs", "Carnet d'adresses. Le fournisseur habituel de chaque produit se choisit dans cette liste.", T)
    heads = [("Fournisseur", 18), ("Contact", 22), ("Téléphone", 16), ("Jour de commande", 16), ("Jour de livraison", 16), ("Notes", 30)]
    for i, (h, w) in enumerate(heads, start=1):
        header(ws, f"{L(i)}4", h, T)
        ws.column_dimensions[L(i)].width = w
    for k in range(10):
        r = 5 + k
        s = SUPPLIERS[k] if k < len(SUPPLIERS) else (None,) * 5
        for j, v in enumerate((s[0], s[1], s[2], s[3], s[4], None)):
            c = inp(ws, f"{L(j + 1)}{r}", v or None, example=True); c.alignment = LEFT
    dv = DataValidation(type="list", formula1='"lundi,mardi,mercredi,jeudi,vendredi,samedi,dimanche"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("D5:E14")
    ws.freeze_panes = "A5"


# ---------------------------------------------------------------- Mois
WEEK_LABELS = ["Date cmd", "Stock", "Commandé", "Reçu"]
TOTAL_LABELS = ["Commandé", "Reçu", "Écart", "Dernier stock"]


def month_cols():
    """Column letters: weeks[w][label], totals[label], spend."""
    weeks, col = {}, 4
    for w in range(4):
        weeks[w] = {lab: L(col + j) for j, lab in enumerate(WEEK_LABELS)}
        col += 4
    totals = {lab: L(col + j) for j, lab in enumerate(TOTAL_LABELS)}
    return weeks, totals, L(col + len(TOTAL_LABELS))


def sheet_month(wb, name):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    title(ws, f"{name} 2026", "Date cmd = jour où tu passes la commande · Stock = ce qu'il reste · Commandé · Reçu = ce qui est arrivé", T)
    weeks, totals, spend = month_cols()
    for c, h, w in (("A", "Produit", 24), ("B", "Unité", 7), ("C", "Mini", 7)):
        header(ws, f"{c}4", h, T, span_to=f"{c}5")
        ws.column_dimensions[c].width = w
    for w in range(4):
        first, last = weeks[w]["Date cmd"], weeks[w]["Reçu"]
        header(ws, f"{first}4", f"Semaine {w + 1}", T, span_to=f"{last}4")
        for lab, letter in weeks[w].items():
            header(ws, f"{letter}5", lab, T, soft=True)
            ws.column_dimensions[letter].width = 10.5 if lab != "Date cmd" else 11
    header(ws, f"{totals['Commandé']}4", "Total du mois", T, span_to=f"{totals['Dernier stock']}4")
    for lab, letter in totals.items():
        header(ws, f"{letter}5", lab, T, soft=True)
        ws.column_dimensions[letter].width = 11
    header(ws, f"{spend}4", "Dépense (€)", T, span_to=f"{spend}5")
    ws.column_dimensions[spend].width = 12
    for z, r in enumerate(zone_header_rows("month")):
        zone_bar(ws, r, spend, ZONES[z])
    for z, k, pr, mr in ALL:
        p = example_product(z, k)
        cell(ws, f"A{mr}", f'=IF(Produits!$A${pr}="","",Produits!$A${pr})', bold=True)
        cell(ws, f"B{mr}", f'=IF(Produits!$A${pr}="","",Produits!$C${pr})', color="555555", align="center")
        cell(ws, f"C{mr}", f'=IF(Produits!$A${pr}="","",Produits!$D${pr})', color="555555", align="center")
        for w in range(4):
            day = WEEK_DATES.get(name, [])
            date = f"=DATE(2026,{MONTHS.index(name) + 1},{day[w]})" if w < len(day) and p else None
            inp(ws, f"{weeks[w]['Date cmd']}{mr}", date, example=True, fmt="DD/MM/YYYY")
            s, c, rc = week_values(name, w, p)
            inp(ws, f"{weeks[w]['Stock']}{mr}", s, example=True)
            inp(ws, f"{weeks[w]['Commandé']}{mr}", c, example=True)
            inp(ws, f"{weeks[w]['Reçu']}{mr}", rc, example=True)
        cmds = [f"{weeks[w]['Commandé']}{mr}" for w in range(4)]
        recs = [f"{weeks[w]['Reçu']}{mr}" for w in range(4)]
        stocks = [f"{weeks[w]['Stock']}{mr}" for w in range(4)]
        calc(ws, f"{totals['Commandé']}{mr}", non_blank_sum(cmds), bold=True)
        calc(ws, f"{totals['Reçu']}{mr}", non_blank_sum(recs), bold=True)
        a, b = f"{totals['Commandé']}{mr}", f"{totals['Reçu']}{mr}"
        calc(ws, f"{totals['Écart']}{mr}", f'=IF(OR({a}="",{b}=""),"",{b}-{a})', bold=True)
        calc(ws, f"{totals['Dernier stock']}{mr}", last_filled(stocks), bold=True)
        calc(ws, f"{spend}{mr}", f'=IF({b}="","",{b}*N(Produits!$E${pr}))', fmt='#,##0.00 "€"')
    tr = LAST_MR + 2
    header(ws, f"A{tr}", "Total des dépenses du mois", T, span_to=f"C{tr}")
    calc(ws, f"{spend}{tr}", f"=SUM({spend}7:{spend}{LAST_MR})", fmt='#,##0.00 "€"', bold=True, fill=T.soft)
    # conditional formats (all on-sheet, Google Sheets compatible)
    e = totals["Écart"]
    negative_red(ws, f"{e}7:{e}{LAST_MR}", f"{e}7")
    for w in range(4):
        s, c, rc = weeks[w]["Stock"], weeks[w]["Commandé"], weeks[w]["Reçu"]
        ws.conditional_formatting.add(f"{s}7:{s}{LAST_MR}",
                                      FormulaRule(formula=[f"AND(ISNUMBER({s}7),ISNUMBER($C7),{s}7<=$C7)"], fill=ORANGE_FILL))
        ws.conditional_formatting.add(f"{rc}7:{rc}{LAST_MR}",
                                      FormulaRule(formula=[f"AND(ISNUMBER({rc}7),ISNUMBER({c}7),{rc}7<{c}7)"], fill=RED_FILL))
        nonneg_validation(ws, f"{s}7:{rc}{LAST_MR}")
    ws.freeze_panes = "D7"
    return ws


# ---------------------------------------------------------------- Tableau de bord
def sheet_dashboard(wb):
    db = wb.create_sheet("Tableau de bord")
    db.sheet_view.showGridLines = False
    title(db, "Tableau de bord", "Choisis le mois dans la case jaune — tout se recalcule. Rien d'autre à remplir ici.", T)
    cell(db, "A3", "Mois :", bold=True, border=False, align="right")
    c = inp(db, "B3", "Octobre"); c.font = Font(name=FONT, bold=True, size=12)
    dv = DataValidation(type="list", formula1='"' + ",".join(MONTHS) + '"', allow_blank=False)
    db.add_data_validation(dv); dv.add("B3")
    # month list + index (hidden column Z)
    for i, m in enumerate(MONTHS):
        db[f"Z{5 + i}"] = m
    db["Z3"] = "=MATCH($B$3,$Z$5:$Z$16,0)"
    db.column_dimensions["Z"].hidden = True
    weeks, totals, spend = month_cols()

    def pick(col, mr):
        # an empty input cell reached through CHOOSE reads as 0 in Excel: keep only real numbers
        refs = ",".join(f"'{m}'!{col}{mr}" for m in MONTHS)
        return f'=IF(ISNUMBER(CHOOSE($Z$3,{refs})),CHOOSE($Z$3,{refs}),"")'

    # hidden helper block: AA.. per product slot
    helper_cols = ["Produit", "S1", "S2", "S3", "S4", "Commandé", "Reçu", "Écart", "Dernier", "Mini", "Dépense", "SousMini", "Tri"]
    HC = {lab: L(27 + j) for j, lab in enumerate(helper_cols)}
    for lab, letter in HC.items():
        header(db, f"{letter}4", lab, T, soft=True)
        db.column_dimensions[letter].hidden = True
    for n, (z, k, pr, mr) in enumerate(ALL):
        r = 5 + n
        calc(db, f"{HC['Produit']}{r}", f'=IF(Produits!$A${pr}="","",Produits!$A${pr})')
        for w in range(4):
            calc(db, f"{HC[f'S{w + 1}']}{r}", pick(weeks[w]["Stock"], mr))
        calc(db, f"{HC['Commandé']}{r}", pick(totals["Commandé"], mr))
        calc(db, f"{HC['Reçu']}{r}", pick(totals["Reçu"], mr))
        calc(db, f"{HC['Écart']}{r}", pick(totals["Écart"], mr))
        calc(db, f"{HC['Dernier']}{r}", pick(totals["Dernier stock"], mr))
        calc(db, f"{HC['Mini']}{r}", f'=IF(Produits!$A${pr}="","",Produits!$D${pr})')
        calc(db, f"{HC['Dépense']}{r}", pick(spend, mr))
        d, m = f"{HC['Dernier']}{r}", f"{HC['Mini']}{r}"
        calc(db, f"{HC['SousMini']}{r}", f"=IF(AND(ISNUMBER({d}),ISNUMBER({m})),IF({d}<={m},1,0),0)")
        calc(db, f"{HC['Tri']}{r}", f'=IF(OR({HC["Produit"]}{r}="",{HC["Dépense"]}{r}=""),-1,{HC["Dépense"]}{r}+ROW()/1000000)')
    hl = 4 + len(ALL)
    # KPIs
    prod_ranges = "+".join(f"COUNTA(Produits!A{6 + z * (PER_ZONE + 1)}:A{5 + z * (PER_ZONE + 1) + PER_ZONE})" for z in range(len(ZONES)))
    kpis = [("Dépensé ce mois", f"=SUM({HC['Dépense']}5:{HC['Dépense']}{hl})", '#,##0.00 "€"'),
            ("Reçu incomplet", f'=COUNTIF({HC["Écart"]}5:{HC["Écart"]}{hl},"<0")', "0"),
            ("Sous le minimum", f"=SUM({HC['SousMini']}5:{HC['SousMini']}{hl})", "0"),
            ("Produits suivis", f"={prod_ranges}", "0")]
    for k, (lab, f, fmt) in enumerate(kpis):
        c1, c2 = L(1 + 2 * k), L(2 + 2 * k)
        db.merge_cells(f"{c1}5:{c2}5"); db.merge_cells(f"{c1}6:{c2}7")
        cell(db, f"{c1}5", lab, bold=True, color="FFFFFF", fill=T.fill, align="center", size=9)
        db[f"{c2}5"].fill = T.fill
        cell(db, f"{c1}6", f, bold=True, size=18, align="center", fmt=fmt, fill=T.softer)
    db.conditional_formatting.add("C6", FormulaRule(formula=["C6>0"], fill=RED_FILL))
    db.conditional_formatting.add("E6", FormulaRule(formula=["E6>0"], fill=ORANGE_FILL))
    for c in "ABCDEFGH":
        db.column_dimensions[c].width = 14
    # Top 10
    cell(db, "A10", "Top 10 des dépenses du mois", bold=True, size=12, color=T.main, border=False)
    for j, h in enumerate(("Rang", "Produit", "", "Reçu", "Dépense (€)")):
        if h:
            header(db, f"{L(1 + j)}11", h, T)
    db.merge_cells("B11:C11")
    for k in range(10):
        r = 12 + k
        big = f"LARGE(${HC['Tri']}$5:${HC['Tri']}${hl},{k + 1})"
        idx = f"MATCH({big},${HC['Tri']}$5:${HC['Tri']}${hl},0)"
        cell(db, f"A{r}", k + 1, align="center")
        db.merge_cells(f"B{r}:C{r}")
        c = calc(db, f"B{r}", f'=IF({big}<=0,"",INDEX(${HC["Produit"]}$5:${HC["Produit"]}${hl},{idx}))'); c.alignment = LEFT
        calc(db, f"D{r}", f'=IF(B{r}="","",INDEX(${HC["Reçu"]}$5:${HC["Reçu"]}${hl},{idx}))')
        calc(db, f"E{r}", f'=IF(B{r}="","",INDEX(${HC["Dépense"]}$5:${HC["Dépense"]}${hl},{idx}))', fmt='#,##0.00 "€"')
    ch = BarChart(); ch.type = "bar"; ch.title = "Dépenses du mois (€)"; ch.style = 2; ch.legend = None
    ch.add_data(Reference(db, min_col=5, min_row=11, max_row=21), titles_from_data=True)
    ch.set_categories(Reference(db, min_col=2, min_row=12, max_row=21))
    ch.x_axis.scaling.orientation = "maxMin"; ch.height, ch.width = 7.5, 14
    ch.series[0].graphicalProperties.solidFill = T.main
    db.add_chart(ch, "G10")
    # Trends
    tr0 = 24
    cell(db, f"A{tr0}", "Tendance du stock et alertes", bold=True, size=12, color=T.main, border=False)
    th = [("Produit", 2), ("S1", 1), ("S2", 1), ("S3", 1), ("S4", 1), ("Évol.", 1), ("Dernier stock", 1), ("Mini", 1), ("Alerte", 2)]
    col = 1
    cols = {}
    for lab, span in th:
        end = col + span - 1
        header(db, f"{L(col)}{tr0 + 1}", lab, T, span_to=f"{L(end)}{tr0 + 1}" if span > 1 else None)
        cols[lab] = (L(col), L(end))
        col = end + 1
    for n in range(len(ALL)):
        r = tr0 + 2 + n
        h = 5 + n
        db.merge_cells(f"A{r}:B{r}")
        c = cell(db, f"A{r}", f"={HC['Produit']}{h}", bold=True); c.alignment = LEFT
        srefs = [f"{HC[f'S{w + 1}']}{h}" for w in range(4)]
        mx = f"MAX({','.join(srefs)})"
        for w in range(4):
            c = calc(db, f"{cols[f'S{w + 1}'][0]}{r}",
                     f'=IF(OR({srefs[w]}="",{mx}=0),"",REPT("█",MAX(1,ROUND({srefs[w]}/{mx}*6,0))))')
            c.font = Font(name=FONT, size=8, color=T.main); c.alignment = Alignment(horizontal="left", vertical="bottom")
        first_v = f'IF({srefs[0]}<>"",{srefs[0]},IF({srefs[1]}<>"",{srefs[1]},IF({srefs[2]}<>"",{srefs[2]},{srefs[3]})))'
        last_v = last_filled(srefs)[1:]
        calc(db, f"{cols['Évol.'][0]}{r}", f'=IF(COUNT({",".join(srefs)})<2,"",IF({last_v}>{first_v},"▲",IF({last_v}<{first_v},"▼","=")))', bold=True)
        calc(db, f"{cols['Dernier stock'][0]}{r}", f"={HC['Dernier']}{h}", bold=True)
        calc(db, f"{cols['Mini'][0]}{r}", f"={HC['Mini']}{h}")
        db.merge_cells(f"{cols['Alerte'][0]}{r}:{cols['Alerte'][1]}{r}")
        calc(db, f"{cols['Alerte'][0]}{r}", f'=IF({HC["SousMini"]}{h}=1,"À COMMANDER",IF({HC["Produit"]}{h}="","",""))', bold=True)
    last_tr = tr0 + 1 + len(ALL)
    ev = cols["Évol."][0]
    db.conditional_formatting.add(f"{ev}{tr0 + 2}:{ev}{last_tr}", FormulaRule(formula=[f'{ev}{tr0 + 2}="▼"'], font=Font(color="C0392B"), fill=RED_FILL))
    db.conditional_formatting.add(f"{ev}{tr0 + 2}:{ev}{last_tr}", FormulaRule(formula=[f'{ev}{tr0 + 2}="▲"'], font=Font(color="1E7B34"), fill=GREEN_FILL))
    al = cols["Alerte"][0]
    db.conditional_formatting.add(f"{al}{tr0 + 2}:{al}{last_tr}", FormulaRule(formula=[f'{al}{tr0 + 2}="À COMMANDER"'], fill=ORANGE_FILL))
    for lab in ("S1", "S2", "S3", "S4"):
        db.column_dimensions[cols[lab][0]].width = 9
    db.column_dimensions[cols["Évol."][0]].width = 7
    db.freeze_panes = "A5"
    return db


# ---------------------------------------------------------------- Année
def sheet_year(wb):
    ws = wb.create_sheet("Année 2026")
    ws.sheet_view.showGridLines = False
    title(ws, "Année 2026", "Calculé à partir des 12 onglets mensuels — rien à remplir", T)
    weeks, totals, spend = month_cols()
    blocks = [("Quantités reçues par mois", totals["Reçu"], "General", True),
              ("Dépenses par mois (€)", spend, '#,##0.00 "€"', False),
              ("Stock en fin de mois (dernier comptage)", totals["Dernier stock"], "General", False)]
    r = 4
    for label, src_col, fmt, stats in blocks:
        cell(ws, f"A{r}", label, bold=True, size=12, color=T.main, border=False)
        r += 1
        heads = [("Produit", 24), ("Unité", 7)] + [(m, 8) for m in MONTH_SHORT] + [("Total", 10)]
        if stats:
            heads += [("Moyenne / mois", 15), ("Plus gros mois", 15)]
        for i, (h, w) in enumerate(heads, start=1):
            header(ws, f"{L(i)}{r}", h, T)
            ws.column_dimensions[L(i)].width = max(ws.column_dimensions[L(i)].width or 0, w)
        hr = r
        r += 1
        first_data = r
        for z, k, pr, mr in ALL:
            cell(ws, f"A{r}", f'=IF(Produits!$A${pr}="","",Produits!$A${pr})', bold=True)
            cell(ws, f"B{r}", f'=IF(Produits!$A${pr}="","",Produits!$C${pr})', color="555555", align="center")
            for m, mname in enumerate(MONTHS):
                src = f"'{mname}'!{src_col}{mr}"
                calc(ws, f"{L(3 + m)}{r}", f'=IF(N({src})=0,"",N({src}))', fmt=fmt)
            rng = f"C{r}:N{r}"
            if label.startswith("Stock"):
                calc(ws, f"O{r}", last_filled([f"{L(3 + m)}{r}" for m in range(12)]), bold=True)
            else:
                calc(ws, f"O{r}", f'=IF(COUNT({rng})=0,"",SUM({rng}))', fmt=fmt, bold=True)
            if stats:
                calc(ws, f"P{r}", f'=IF(COUNT({rng})=0,"",ROUND(AVERAGE({rng}),1))', bold=True)
                calc(ws, f"Q{r}", f'=IF(COUNT({rng})=0,"",MAX({rng}))')
            r += 1
        if label.startswith("Dépenses"):
            header(ws, f"A{r}", "Total", T, span_to=f"B{r}")
            for m in range(13):
                col = L(3 + m)
                calc(ws, f"{col}{r}", f"=SUM({col}{first_data}:{col}{r - 1})", fmt=fmt, bold=True, fill=T.soft)
            r += 1
        if label.startswith("Stock"):
            ws[f"O{hr}"].value = "Dernier"
        r += 2
    ws.freeze_panes = "C6"


# ---------------------------------------------------------------- Registre de réception
def sheet_register(wb):
    ws = wb.create_sheet("Réceptions")
    ws.sheet_view.showGridLines = False
    title(ws, "Registre de réception", "Une ligne par produit reçu — remplace le cahier. Les cases vides sont permises (pas de lot, pas de DLC).", T)
    heads = [("Reçu le", 12), ("Produit", 24), ("Fournisseur", 16), ("Quantité", 10), ("Unité", 8), ("N° de lot", 18),
             ("DLC", 12), ("Jours restants", 13), ("État", 13), ("Terminé ?", 10)]
    for i, (h, w) in enumerate(heads, start=1):
        header(ws, f"{L(i)}4", h, T)
        ws.column_dimensions[L(i)].width = w
    rows = 400
    for k in range(rows):
        r = 5 + k
        ex = REGISTER_EXAMPLES[k] if k < len(REGISTER_EXAMPLES) else (None,) * 6
        inp(ws, f"A{r}", ex[0], example=True, fmt="DD/MM/YYYY")
        c = inp(ws, f"B{r}", ex[1], example=True); c.alignment = LEFT
        c = inp(ws, f"C{r}", ex[2], example=True); c.alignment = LEFT
        inp(ws, f"D{r}", ex[3], example=True)
        calc(ws, f"E{r}", f'=IF(B{r}="","",IFERROR(INDEX({UNIT_LIST},MATCH(B{r},{PRODUCT_LIST},0)),""))')
        c = inp(ws, f"F{r}", ex[4], example=True); c.alignment = LEFT
        inp(ws, f"G{r}", ex[5], example=True, fmt="DD/MM/YYYY")
        calc(ws, f"H{r}", f'=IF(G{r}="","",G{r}-TODAY())')
        calc(ws, f"I{r}", f'=IF(G{r}="","",IF(J{r}="Oui","Terminé",IF(H{r}<0,"PÉRIMÉ",IF(H{r}<=1,"AUJOURD\'HUI",IF(H{r}<=3,"BIENTÔT","OK")))))')
        inp(ws, f"J{r}", None)
    last = 4 + rows
    dv = DataValidation(type="list", formula1="=" + PRODUCT_LIST, allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv); dv.add(f"B5:B{last}")
    dv2 = DataValidation(type="list", formula1="=Fournisseurs!$A$5:$A$14", allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv2); dv2.add(f"C5:C{last}")
    dv3 = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
    ws.add_data_validation(dv3); dv3.add(f"J5:J{last}")
    nonneg_validation(ws, f"D5:D{last}")
    grey = DifferentialStyle(font=Font(color="8C8C8C"), fill=GREY_FILL)
    ws.conditional_formatting.add(f"A5:J{last}", Rule(type="expression", dxf=grey, formula=['$J5="Oui"']))
    for text, fill in (("PÉRIMÉ", RED_FILL), ("AUJOURD'HUI", RED_FILL), ("BIENTÔT", ORANGE_FILL), ("OK", GREEN_FILL)):
        ws.conditional_formatting.add(f"I5:I{last}", FormulaRule(formula=[f'I5="{text}"'], fill=fill))
    ws["L4"] = "Fournisseur : choisis dans la liste ou écris-en un autre si tu as acheté ailleurs."
    ws["L4"].font = Font(name=FONT, italic=True, size=9, color="666666")
    ws.freeze_panes = "B5"


# ---------------------------------------------------------------- build
def build():
    wb = Workbook()
    guide(wb, T, "Suivi des stocks — version pour Lele (v1)",
          "Ton suivi sur l'année, avec tout ce que tu as demandé : les produits rangés par zone (réserve sèche, "
          "chambre froide positive, chambre froide négative, ménager), la date de commande chaque semaine, le "
          "tableau de bord avec les tendances, et le registre de réception qui remplace ton cahier.",
          ["1. Onglet Produits : écris chaque produit dans sa zone, avec son fournisseur habituel, l'unité, le stock "
           "minimum et le prix unitaire (une seule fois — à changer seulement si le tarif change).",
           "2. Chaque semaine, dans l'onglet du mois : la date de commande, le stock compté, ce que tu commandes, "
           "puis ce que tu reçois à la livraison. Rouge = reçu moins que commandé, orange = sous le minimum.",
           "3. À chaque livraison : une ligne par produit dans Réceptions (date, produit, fournisseur, quantité, "
           "n° de lot, DLC). Le fournisseur peut être différent de l'habituel.",
           "4. Tableau de bord : choisis le mois, tout se calcule. Année 2026 : le récap mois par mois.",
           "5. Les chiffres en bleu sont des exemples : efface-les avant de commencer."],
          ["Tableau de bord — dépenses, livraisons incomplètes, produits sous le minimum, tendances, top 10.",
           "Produits — ta liste par zone, avec minimum et prix.", "Fournisseurs — carnet d'adresses.",
           "Janvier à Décembre — la saisie de chaque semaine.", "Année 2026 — reçu, dépenses et stock de fin de mois.",
           "Réceptions — le registre qui remplace le cahier."])
    sheet_products(wb)
    sheet_suppliers(wb)
    for m in MONTHS:
        sheet_month(wb, m)
    sheet_year(wb)
    sheet_register(wb)
    sheet_dashboard(wb)
    order = ["Mode d'emploi", "Tableau de bord", "Produits", "Fournisseurs", "Réceptions"] + MONTHS + ["Année 2026"]
    wb._sheets = [wb[n] for n in order]
    save(wb, "30_version_lele_v1.xlsx")
    print(wb.sheetnames)


if __name__ == "__main__":
    build()
